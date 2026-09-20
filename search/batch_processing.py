"""Batch import/export for organization payroll classification.

The batch workflow deliberately reuses :class:`OrganizationSearchService` for
every input value.  It therefore has the same exact-first and conservative
fuzzy guarantees as the single-record API and never infers payroll from the
management field.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
import re
from typing import Any, Sequence
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from preprocessing.normalize import to_search_key
from search.organization_search import OrganizationSearchService


MAX_BATCH_ROWS = 5_000
MAX_EXPANDED_OFFICE_BYTES = 100 * 1024 * 1024
MAX_OFFICE_ARCHIVE_ENTRIES = 10_000

PAID_STATUSES = {"Do BCA trả lương", "Do BQP trả lương"}
NOT_PAID_STATUS = "Không do BCA/BQP trả lương"

_HEADER_KEYS = {
    "ten don vi",
    "ten to chuc",
    "don vi",
    "to chuc",
    "organization name",
    "organization_name",
    "org name",
}
_GENERIC_HEADER_KEYS = _HEADER_KEYS | {"ten", "name"}
_LIST_PREFIX = re.compile(r"^\s*(?:(?:[-•●▪◦])|(?:\d{1,5}[.)]))\s+")


class BatchProcessingError(ValueError):
    """Raised when an uploaded Office document cannot be safely processed."""


@dataclass(frozen=True, slots=True)
class InputOrganization:
    source: str
    source_row: int
    organization_name: str


@dataclass(frozen=True, slots=True)
class BatchRow:
    source: str
    source_row: int
    input_name: str
    category: str
    result: dict[str, Any]


@dataclass(frozen=True, slots=True)
class BatchArtifact:
    archive: bytes
    archive_name: str
    input_count: int
    paid_count: int
    not_paid_count: int
    unresolved_count: int


def process_batch_file(
    filename: str,
    content: bytes,
    service: OrganizationSearchService,
    column_name: str | None = None,
) -> BatchArtifact:
    """Read a Word/Excel list, classify rows, and return two reports in a ZIP."""

    if not filename:
        raise BatchProcessingError("Tên file không hợp lệ.")
    if not content:
        raise BatchProcessingError("File tải lên đang trống.")

    extension = Path(filename).suffix.lower()
    if extension not in {".xlsx", ".xls", ".docx"}:
        raise BatchProcessingError(
            "Chỉ hỗ trợ file Excel .xlsx/.xls hoặc Word .docx."
        )
    if extension in {".xlsx", ".docx"}:
        _validate_office_archive(content)

    if extension in {".xlsx", ".xls"}:
        inputs = _read_excel(content, extension, column_name)
        output_extension = ".xlsx"
    else:
        inputs = _read_word(content, column_name)
        output_extension = ".docx"

    if not inputs:
        raise BatchProcessingError(
            "Không tìm thấy tên đơn vị nào trong file. Hãy kiểm tra cột hoặc danh sách đầu vào."
        )
    if len(inputs) > MAX_BATCH_ROWS:
        raise BatchProcessingError(
            f"File có {len(inputs):,} dòng; giới hạn mỗi lần xử lý là {MAX_BATCH_ROWS:,} dòng."
        )

    # Reuse the complete search response for repeated names. Large operational
    # lists often contain duplicates, and fuzzy matching is intentionally the
    # most expensive stage of the resolver.
    result_cache: dict[str, dict[str, Any]] = {}
    classified: list[BatchRow] = []
    for item in inputs:
        result = result_cache.get(item.organization_name)
        if result is None:
            result = service.search_organization(organization_name=item.organization_name)
            result_cache[item.organization_name] = result
        classified.append(_classify_result(item, result))
    rows = tuple(classified)
    paid = tuple(row for row in rows if row.category == "PAID")
    not_paid = tuple(row for row in rows if row.category == "NOT_PAID")
    unresolved = tuple(row for row in rows if row.category == "UNRESOLVED")

    if output_extension == ".xlsx":
        paid_bytes = _build_excel_report(
            "Danh sách đơn vị được trả lương", paid, (), "Được trả lương"
        )
        not_paid_bytes = _build_excel_report(
            "Danh sách đơn vị không được trả lương",
            not_paid,
            unresolved,
            "Không được trả lương",
        )
    else:
        paid_bytes = _build_word_report(
            "Danh sách đơn vị được trả lương", paid, ()
        )
        not_paid_bytes = _build_word_report(
            "Danh sách đơn vị không được trả lương", not_paid, unresolved
        )

    stem = _safe_stem(filename)
    paid_name = f"{stem}_duoc_tra_luong{output_extension}"
    not_paid_name = f"{stem}_khong_duoc_tra_luong{output_extension}"
    archive_name = f"{stem}_ket_qua_tra_luong.zip"
    archive_buffer = BytesIO()
    with ZipFile(archive_buffer, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr(paid_name, paid_bytes)
        archive.writestr(not_paid_name, not_paid_bytes)

    return BatchArtifact(
        archive=archive_buffer.getvalue(),
        archive_name=archive_name,
        input_count=len(rows),
        paid_count=len(paid),
        not_paid_count=len(not_paid),
        unresolved_count=len(unresolved),
    )


def _validate_office_archive(content: bytes) -> None:
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_OFFICE_ARCHIVE_ENTRIES:
                raise BatchProcessingError("File Office có cấu trúc quá lớn.")
            expanded_size = sum(entry.file_size for entry in entries)
            if expanded_size > MAX_EXPANDED_OFFICE_BYTES:
                raise BatchProcessingError("Dung lượng giải nén của file Office vượt giới hạn.")
    except BadZipFile as error:
        raise BatchProcessingError("File Office bị hỏng hoặc không đúng định dạng.") from error


def _read_excel(
    content: bytes, extension: str, column_name: str | None
) -> tuple[InputOrganization, ...]:
    if extension == ".xlsx":
        return _read_xlsx(content, column_name)
    return _read_xls(content, column_name)


def _read_xlsx(
    content: bytes, column_name: str | None
) -> tuple[InputOrganization, ...]:
    try:
        from openpyxl import load_workbook

        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as error:
        raise BatchProcessingError("Không thể đọc file Excel .xlsx.") from error

    found: list[InputOrganization] = []
    errors: list[BatchProcessingError] = []
    try:
        for worksheet in workbook.worksheets:
            if worksheet.max_row > MAX_BATCH_ROWS + 50:
                raise BatchProcessingError(
                    f"Sheet “{worksheet.title}” có phạm vi {worksheet.max_row:,} dòng; "
                    f"giới hạn cho phép là {MAX_BATCH_ROWS:,} dòng dữ liệu."
                )
            if worksheet.max_column > 256:
                raise BatchProcessingError(
                    f"Sheet “{worksheet.title}” có quá nhiều cột để xử lý an toàn."
                )
            rows = [tuple(row) for row in worksheet.iter_rows(values_only=True)]
            try:
                found.extend(_extract_tabular_rows(rows, worksheet.title, column_name))
            except BatchProcessingError as error:
                errors.append(error)
            _ensure_row_limit(found)
    finally:
        workbook.close()
    if not found and errors:
        raise errors[0]
    return tuple(found)


def _read_xls(
    content: bytes, column_name: str | None
) -> tuple[InputOrganization, ...]:
    try:
        import pandas as pd

        sheets = pd.read_excel(BytesIO(content), sheet_name=None, header=None, dtype=object)
    except Exception as error:
        raise BatchProcessingError(
            "Không thể đọc file Excel .xls. Hãy kiểm tra file hoặc lưu lại dưới dạng .xlsx."
        ) from error

    found: list[InputOrganization] = []
    errors: list[BatchProcessingError] = []
    for sheet_name, frame in sheets.items():
        rows = [
            tuple(row)
            for row in frame.where(frame.notna(), None).itertuples(
                index=False, name=None
            )
        ]
        try:
            found.extend(_extract_tabular_rows(rows, str(sheet_name), column_name))
        except BatchProcessingError as error:
            errors.append(error)
        _ensure_row_limit(found)
    if not found and errors:
        raise errors[0]
    return tuple(found)


def _extract_tabular_rows(
    rows: Sequence[Sequence[Any]], source: str, column_name: str | None
) -> list[InputOrganization]:
    if not rows:
        return []

    requested_key = _header_key(column_name) if column_name else None
    header_row_index: int | None = None
    column_index: int | None = None

    for row_index, row in enumerate(rows[:25]):
        for candidate_index, value in enumerate(row):
            key = _header_key(value)
            if requested_key and key == requested_key:
                header_row_index, column_index = row_index, candidate_index
                break
            if not requested_key and key in _HEADER_KEYS:
                header_row_index, column_index = row_index, candidate_index
                break
        if column_index is not None:
            break

    if column_index is None:
        populated_columns = {
            column_index
            for row in rows
            for column_index, value in enumerate(row)
            if _clean_cell(value)
        }
        if requested_key:
            raise BatchProcessingError(
                f"Không tìm thấy cột “{column_name}” trong sheet “{source}”."
            )
        if len(populated_columns) != 1:
            raise BatchProcessingError(
                f"Không xác định được cột tên đơn vị trong sheet “{source}”. "
                "Hãy dùng tiêu đề “Tên đơn vị” hoặc nhập rõ tên cột trên giao diện."
            )
        column_index = next(iter(populated_columns))
        first_value = _clean_cell(
            rows[0][column_index] if column_index < len(rows[0]) else None
        )
        header_row_index = 0 if _header_key(first_value) in _GENERIC_HEADER_KEYS else -1

    start = (header_row_index or 0) + 1 if header_row_index is not None else 0
    extracted: list[InputOrganization] = []
    for zero_based_index, row in enumerate(rows[start:], start=start):
        value = row[column_index] if column_index < len(row) else None
        name = _clean_cell(value)
        if name:
            extracted.append(
                InputOrganization(
                    source=source,
                    source_row=zero_based_index + 1,
                    organization_name=name,
                )
            )
    return extracted


def _read_word(content: bytes, column_name: str | None) -> tuple[InputOrganization, ...]:
    try:
        from docx import Document

        document = Document(BytesIO(content))
    except Exception as error:
        raise BatchProcessingError("Không thể đọc file Word .docx.") from error

    found: list[InputOrganization] = []
    errors: list[BatchProcessingError] = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = [tuple(cell.text for cell in row.cells) for row in table.rows]
        try:
            found.extend(_extract_tabular_rows(rows, f"Bảng {table_index}", column_name))
        except BatchProcessingError as error:
            errors.append(error)
        _ensure_row_limit(found)

    if found:
        return tuple(found)
    if column_name:
        if errors:
            raise errors[0]
        raise BatchProcessingError(
            f"Không tìm thấy cột “{column_name}” trong các bảng của file Word."
        )

    paragraph_row = 0
    for paragraph in document.paragraphs:
        paragraph_row += 1
        style_name = (paragraph.style.name if paragraph.style else "").lower()
        if style_name.startswith("title") or style_name.startswith("heading"):
            continue
        name = _LIST_PREFIX.sub("", paragraph.text).strip()
        if not name or _header_key(name) in _GENERIC_HEADER_KEYS:
            continue
        found.append(
            InputOrganization(
                source="Đoạn văn",
                source_row=paragraph_row,
                organization_name=name,
            )
        )
        _ensure_row_limit(found)
    return tuple(found)


def _ensure_row_limit(rows: Sequence[Any]) -> None:
    if len(rows) > MAX_BATCH_ROWS:
        raise BatchProcessingError(
            f"File vượt giới hạn {MAX_BATCH_ROWS:,} tên đơn vị mỗi lần xử lý."
        )


def _classify_result(item: InputOrganization, result: dict[str, Any]) -> BatchRow:
    payroll_status = result.get("payroll_status")
    if result.get("organization_id") and payroll_status in PAID_STATUSES:
        category = "PAID"
    elif result.get("organization_id") and payroll_status == NOT_PAID_STATUS:
        category = "NOT_PAID"
    else:
        category = "UNRESOLVED"
    return BatchRow(
        source=item.source,
        source_row=item.source_row,
        input_name=item.organization_name,
        category=category,
        result=result,
    )


def _build_excel_report(
    title: str,
    rows: Sequence[BatchRow],
    unresolved: Sequence[BatchRow],
    sheet_name: str,
) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.sheet_view.showGridLines = False

    generated_at = datetime.now(timezone.utc).astimezone().strftime("%d/%m/%Y %H:%M")
    worksheet.merge_cells("A1:J1")
    worksheet["A1"] = _excel_text(title.upper())
    worksheet["A1"].font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    worksheet["A1"].fill = PatternFill("solid", fgColor="163A5F")
    worksheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 30
    worksheet.merge_cells("A2:J2")
    worksheet["A2"] = f"Tổng số: {len(rows):,} đơn vị · Xuất lúc {generated_at}"
    worksheet["A2"].font = Font(name="Arial", size=10, color="5B677A")

    headers = [
        "STT",
        "Tên đơn vị nhập",
        "Tên chính thức",
        "Mã tổ chức",
        "Đơn vị trả lương",
        "Trạng thái trả lương",
        "Phạm vi quản lý",
        "Phương thức khớp",
        "Điểm khớp",
        "Nguồn / dòng",
    ]
    header_row = 4
    for column, value in enumerate(headers, start=1):
        cell = worksheet.cell(header_row, column, value)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="246B92")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    worksheet.row_dimensions[header_row].height = 32

    for index, row in enumerate(rows, start=1):
        result = row.result
        values = [
            index,
            row.input_name,
            result.get("organization_name") or "",
            result.get("organization_id") or "",
            result.get("paying_organization") or "",
            result.get("payroll_status") or "",
            result.get("management") or "",
            result.get("match_status") or "",
            result.get("match_score") if result.get("match_score") is not None else "",
            f"{row.source} / dòng {row.source_row}",
        ]
        excel_row = header_row + index
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(excel_row, column, _excel_text(value))
            cell.font = Font(name="Arial", size=10, color="1D2A38")
            cell.alignment = Alignment(
                horizontal="center" if column in {1, 7, 9} else "left",
                vertical="center",
                wrap_text=True,
            )
        if index % 2 == 0:
            for cell in worksheet[excel_row]:
                cell.fill = PatternFill("solid", fgColor="F3F7FA")

    last_row = header_row + max(len(rows), 1)
    if rows:
        table = Table(displayName="DanhSachKetQua", ref=f"A{header_row}:J{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=False,
            showColumnStripes=False,
        )
        worksheet.add_table(table)
    else:
        worksheet.merge_cells(start_row=5, start_column=1, end_row=5, end_column=10)
        worksheet["A5"] = "Không có đơn vị thuộc nhóm này."
        worksheet["A5"].font = Font(name="Arial", size=10, italic=True, color="6B7280")
        worksheet["A5"].alignment = Alignment(horizontal="center")

    widths = [7, 34, 34, 24, 23, 29, 18, 24, 12, 24]
    for index, width in enumerate(widths, start=1):
        worksheet.column_dimensions[_excel_column(index)].width = width
    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = f"A{header_row}:J{last_row}" if rows else None
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = f"1:{header_row}"

    if unresolved:
        _add_unresolved_excel_sheet(workbook, unresolved)

    border = Border(bottom=Side(style="thin", color="D8E0E8"))
    for row_cells in worksheet.iter_rows(min_row=5, max_row=last_row, min_col=1, max_col=10):
        for cell in row_cells:
            cell.border = border

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _add_unresolved_excel_sheet(workbook: Any, rows: Sequence[BatchRow]) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    worksheet = workbook.create_sheet("Chưa thể kết luận")
    worksheet.sheet_view.showGridLines = False
    headers = ["STT", "Tên đơn vị nhập", "Trạng thái đối chiếu", "Lý do / ứng viên", "Nguồn / dòng"]
    worksheet.append(headers)
    for cell in worksheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="9A6700")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for index, row in enumerate(rows, start=1):
        worksheet.append(
            [
                index,
                _excel_text(row.input_name),
                row.result.get("match_status") or "UNKNOWN",
                _excel_text(_unresolved_detail(row.result)),
                _excel_text(f"{row.source} / dòng {row.source_row}"),
            ]
        )
    table = Table(displayName="DanhSachChuaKetLuan", ref=f"A1:E{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
    worksheet.add_table(table)
    for row_cells in worksheet.iter_rows(min_row=2):
        for column, cell in enumerate(row_cells, start=1):
            cell.font = Font(name="Arial", size=10, color="1D2A38")
            cell.alignment = Alignment(
                horizontal="center" if column == 1 else "left",
                vertical="center",
                wrap_text=True,
            )
    for index, width in enumerate([7, 38, 25, 58, 25], start=1):
        worksheet.column_dimensions[_excel_column(index)].width = width
    worksheet.freeze_panes = "A2"
    worksheet.page_setup.orientation = "landscape"
    worksheet.page_setup.fitToWidth = 1
    worksheet.sheet_properties.pageSetUpPr.fitToPage = True
    worksheet.print_title_rows = "1:1"


def _build_word_report(
    title: str, rows: Sequence[BatchRow], unresolved: Sequence[BatchRow]
) -> bytes:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")

    header = section.header.paragraphs[0]
    header.text = "BCA / BQP ORGANIZATION REGISTRY · KẾT QUẢ ĐỐI CHIẾU HÀNG LOẠT"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in header.runs:
        _set_docx_run(run, 8, "66788A", bold=True)

    title_paragraph = document.add_paragraph()
    title_paragraph.paragraph_format.space_before = Pt(0)
    title_paragraph.paragraph_format.space_after = Pt(3)
    title_run = title_paragraph.add_run(title)
    _set_docx_run(title_run, 20, "163A5F", bold=True)

    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(10)
    meta_run = meta.add_run(
        f"Tổng số: {len(rows):,} đơn vị · Xuất lúc "
        f"{datetime.now(timezone.utc).astimezone().strftime('%d/%m/%Y %H:%M')}"
    )
    _set_docx_run(meta_run, 9, "66788A")

    headers = [
        "STT",
        "Tên đơn vị nhập",
        "Tên chính thức",
        "Đơn vị trả lương",
        "Trạng thái",
        "Quản lý",
    ]
    widths = [500, 2450, 2450, 1500, 1900, 800]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_docx_table_geometry(table, widths, 120)
    _repeat_header(table.rows[0])
    for index, header_text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header_text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(cell, "246B92")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                _set_docx_run(run, 8.5, "FFFFFF", bold=True)

    if rows:
        for index, row in enumerate(rows, start=1):
            result = row.result
            values = [
                str(index),
                row.input_name,
                result.get("organization_name") or "",
                result.get("paying_organization") or "",
                result.get("payroll_status") or "",
                result.get("management") or "",
            ]
            cells = table.add_row().cells
            for column, value in enumerate(values):
                cells[column].text = str(value)
                cells[column].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if index % 2 == 0:
                    _shade_cell(cells[column], "F3F7FA")
                for paragraph in cells[column].paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.alignment = (
                        WD_ALIGN_PARAGRAPH.CENTER
                        if column in {0, 5}
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                    for run in paragraph.runs:
                        _set_docx_run(run, 8.5, "1D2A38")
    else:
        cells = table.add_row().cells
        merged = cells[0]
        for cell in cells[1:]:
            merged = merged.merge(cell)
        merged.text = "Không có đơn vị thuộc nhóm này."
        merged.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in merged.paragraphs[0].runs:
            _set_docx_run(run, 9, "66788A", italic=True)

    _set_docx_table_geometry(table, widths, 120)

    if unresolved:
        heading = document.add_paragraph()
        heading.paragraph_format.space_before = Pt(14)
        heading.paragraph_format.space_after = Pt(7)
        heading_run = heading.add_run("Phụ lục: Các dòng chưa thể kết luận")
        _set_docx_run(heading_run, 13, "9A6700", bold=True)
        note = document.add_paragraph(
            "Các dòng dưới đây không được gán vào nhóm không được trả lương "
            "vì chưa có kết quả định danh duy nhất, an toàn."
        )
        note.paragraph_format.space_after = Pt(7)
        for run in note.runs:
            _set_docx_run(run, 9, "5B677A", italic=True)
        _add_unresolved_word_table(document, unresolved)

    footer = section.footer.paragraphs[0]
    footer.text = (
        "Kết quả được tạo tự động; cần rà soát các dòng trong phụ lục "
        "trước khi sử dụng nghiệp vụ."
    )
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in footer.runs:
        _set_docx_run(run, 8, "66788A", italic=True)

    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _add_unresolved_word_table(document: Any, rows: Sequence[BatchRow]) -> None:
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    headers = ["STT", "Tên đơn vị nhập", "Trạng thái", "Lý do / ứng viên", "Nguồn / dòng"]
    widths = [500, 2450, 1400, 4050, 1200]
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    _set_docx_table_geometry(table, widths, 120)
    _repeat_header(table.rows[0])
    for index, header_text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = header_text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        _shade_cell(cell, "9A6700")
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                _set_docx_run(run, 8.5, "FFFFFF", bold=True)
    for index, row in enumerate(rows, start=1):
        values = [
            str(index),
            row.input_name,
            row.result.get("match_status") or "UNKNOWN",
            _unresolved_detail(row.result),
            f"{row.source} / dòng {row.source_row}",
        ]
        cells = table.add_row().cells
        for column, value in enumerate(values):
            cells[column].text = str(value)
            cells[column].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if index % 2 == 0:
                _shade_cell(cells[column], "FFF7E0")
            for paragraph in cells[column].paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER if column == 0 else WD_ALIGN_PARAGRAPH.LEFT
                )
                for run in paragraph.runs:
                    _set_docx_run(run, 8.5, "1D2A38")
    _set_docx_table_geometry(table, widths, 120)


def _set_docx_run(
    run: Any,
    size: float,
    color: str,
    *,
    bold: bool = False,
    italic: bool = False,
) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = "Calibri"
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def _shade_cell(cell: Any, fill: str) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _repeat_header(row: Any) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _set_docx_table_geometry(table: Any, widths: Sequence[int], indent: int) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    total = sum(widths)
    properties = table._tbl.tblPr
    table_width = properties.find(qn("w:tblW"))
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.append(table_width)
    table_width.set(qn("w:type"), "dxa")
    table_width.set(qn("w:w"), str(total))
    table_indent = properties.find(qn("w:tblInd"))
    if table_indent is None:
        table_indent = OxmlElement("w:tblInd")
        properties.append(table_indent)
    table_indent.set(qn("w:type"), "dxa")
    table_indent.set(qn("w:w"), str(indent))

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            properties = cell._tc.get_or_add_tcPr()
            cell_width = properties.find(qn("w:tcW"))
            if cell_width is None:
                cell_width = OxmlElement("w:tcW")
                properties.append(cell_width)
            cell_width.set(qn("w:type"), "dxa")
            cell_width.set(qn("w:w"), str(widths[index]))
            margins = properties.find(qn("w:tcMar"))
            if margins is None:
                margins = OxmlElement("w:tcMar")
                properties.append(margins)
            for edge, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
                element = margins.find(qn(f"w:{edge}"))
                if element is None:
                    element = OxmlElement(f"w:{edge}")
                    margins.append(element)
                element.set(qn("w:w"), str(value))
                element.set(qn("w:type"), "dxa")


def _unresolved_detail(result: dict[str, Any]) -> str:
    candidates = result.get("candidates") or []
    if candidates:
        names = "; ".join(
            f"{candidate.get('organization_name', '')} ({candidate.get('score', '')}%)"
            for candidate in candidates[:3]
        )
        return f"Cần xác nhận: {names}"
    reason = result.get("reason")
    errors = result.get("errors") or []
    return str(reason or "; ".join(str(item) for item in errors) or "Không có kết quả đủ tin cậy")


def _header_key(value: Any) -> str:
    text = _clean_cell(value)
    if not text:
        return ""
    key = to_search_key(text)
    return key.replace("_", " ").strip()


def _clean_cell(value: Any) -> str:
    if value is None:
        return ""
    text = str(value).replace("\x00", "").strip()
    return text[:512]


def _excel_text(value: Any) -> Any:
    if not isinstance(value, str):
        return value
    # Prevent uploaded values from becoming formulas in generated workbooks.
    return f"'{value}" if value.startswith(("=", "+", "-", "@")) else value


def _excel_column(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)


def _safe_stem(filename: str) -> str:
    stem = Path(filename).stem
    ascii_stem = to_search_key(stem).replace(" ", "_")
    safe = re.sub(r"[^a-z0-9_-]+", "", ascii_stem)[:64]
    return safe or "danh_sach_don_vi"
