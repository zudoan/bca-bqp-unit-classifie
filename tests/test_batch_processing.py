from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
import unittest
from unittest.mock import MagicMock, patch
from zipfile import ZipFile

import pandas as pd
import requests
from docx import Document
from docx.enum.section import WD_ORIENT
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

os.environ["CORS_ORIGINS"] = "https://bca-bqp-frontend.onrender.com/"

from api.main import app
from matching.repository import InMemoryOrganizationRepository
from matching.resolver import MatchingConfig
from search.batch_processing import process_batch_file
from search.organization_search import OrganizationSearchService
from search.pdf_processing import (
    GeminiOcrProvider,
    OcrConfigurationError,
    OcrUpstreamError,
)


ROOT = Path(__file__).resolve().parent.parent


def make_service() -> OrganizationSearchService:
    rows = pd.read_csv(
        ROOT / "data" / "dataset.csv", dtype=str, keep_default_na=False
    ).to_dict("records")
    aliases = pd.read_csv(
        ROOT / "data" / "aliases.csv", dtype=str, keep_default_na=False
    ).to_dict("records")
    return OrganizationSearchService(
        InMemoryOrganizationRepository(rows, aliases), MatchingConfig()
    )


def make_excel_input() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Đơn vị"
    sheet.append(["Tên đơn vị"])
    sheet.append(["Công an tỉnh Thái Bình"])
    sheet.append(["Cục Tổ chức cán bộ"])
    sheet.append(["Đơn vị hoàn toàn không tồn tại 12345"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def make_word_input() -> bytes:
    document = Document()
    document.add_heading("Danh sách đầu vào", level=1)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Tên tổ chức"
    for name in (
        "Công an tỉnh Thái Bình",
        "Cục Tổ chức cán bộ",
        "Đơn vị hoàn toàn không tồn tại 12345",
    ):
        table.add_row().cells[0].text = name
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def make_text_pdf_input() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=595, height=842)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})}
    )
    stream = DecodedStreamObject()
    stream.set_data(
        b"BT /F1 11 Tf 72 770 Td 14 TL "
        b"(Ten don vi) Tj T* "
        b"(Cong an tinh Thai Binh) Tj T* "
        b"(Cuc To chuc can bo) Tj T* "
        b"(Don vi khong ton tai 12345) Tj ET"
    )
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def make_scan_pdf_input() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


class FakeGeminiOcrProvider:
    def __init__(self) -> None:
        self.calls = 0

    def extract_text(self, content: bytes) -> str:
        self.calls += 1
        return (
            "| Tên đơn vị |\n"
            "| --- |\n"
            "| Công an tỉnh Thái Bình |\n"
            "| Cục Tổ chức cán bộ |\n"
            "| Đơn vị không tồn tại 12345 |"
        )


class UnconfiguredOcrProvider:
    def extract_text(self, content: bytes) -> str:
        raise OcrConfigurationError("Gemini OCR chưa được cấu hình cho kiểm thử.")


class BatchProcessingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.service = make_service()

    def test_excel_creates_two_workbooks_and_keeps_unresolved_separate(self) -> None:
        artifact = process_batch_file(
            "input.xlsx", make_excel_input(), self.service
        )

        self.assertEqual(artifact.input_count, 3)
        self.assertEqual(artifact.paid_count, 1)
        self.assertEqual(artifact.not_paid_count, 1)
        self.assertEqual(artifact.unresolved_count, 1)
        self.assertEqual(artifact.input_mode, "excel")

        with ZipFile(BytesIO(artifact.archive)) as archive:
            names = sorted(archive.namelist())
            self.assertEqual(
                names,
                [
                    "input_duoc_tra_luong.xlsx",
                    "input_khong_duoc_tra_luong.xlsx",
                ],
            )
            paid_book = load_workbook(BytesIO(archive.read(names[0])), data_only=True)
            not_paid_book = load_workbook(BytesIO(archive.read(names[1])), data_only=True)

        self.assertIn("Công an tỉnh Thái Bình", paid_book["Được trả lương"][5][1].value)
        self.assertEqual(paid_book["Được trả lương"].freeze_panes, "A5")
        self.assertFalse(paid_book["Được trả lương"].sheet_view.showGridLines)
        self.assertEqual(paid_book["Được trả lương"]["A1"].fill.fgColor.rgb, "00163A5F")
        self.assertEqual(
            not_paid_book.sheetnames,
            ["Không được trả lương", "Chưa thể kết luận"],
        )
        self.assertIn("Cục Tổ chức cán bộ", not_paid_book["Không được trả lương"][5][1].value)
        self.assertIn(
            "Đơn vị hoàn toàn không tồn tại",
            not_paid_book["Chưa thể kết luận"][2][1].value,
        )

    def test_word_creates_two_documents_with_unresolved_appendix(self) -> None:
        artifact = process_batch_file("input.docx", make_word_input(), self.service)

        self.assertEqual(
            (artifact.paid_count, artifact.not_paid_count, artifact.unresolved_count),
            (1, 1, 1),
        )
        with ZipFile(BytesIO(artifact.archive)) as archive:
            names = sorted(archive.namelist())
            self.assertEqual(len(names), 2)
            paid = Document(BytesIO(archive.read("input_duoc_tra_luong.docx")))
            not_paid = Document(BytesIO(archive.read("input_khong_duoc_tra_luong.docx")))

        paid_text = "\n".join(
            cell.text
            for table in paid.tables
            for row in table.rows
            for cell in row.cells
        )
        not_paid_text = "\n".join(
            [paragraph.text for paragraph in not_paid.paragraphs]
            + [cell.text for table in not_paid.tables for row in table.rows for cell in row.cells]
        )
        self.assertIn("Công an tỉnh Thái Bình", paid_text)
        self.assertEqual(paid.sections[0].orientation, WD_ORIENT.LANDSCAPE)
        self.assertEqual(paid.tables[0].style.name, "Table Grid")
        self.assertIsNotNone(paid.tables[0].rows[0]._tr.trPr)
        self.assertIn("tblHeader", paid.tables[0].rows[0]._tr.trPr.xml)
        self.assertIn("Cục Tổ chức cán bộ", not_paid_text)
        self.assertIn("Phụ lục: Các dòng chưa thể kết luận", not_paid_text)

    def test_api_returns_zip_and_summary_headers(self) -> None:
        response = TestClient(app).post(
            "/api/v1/organizations/batch",
            files={
                "file": (
                    "input.xlsx",
                    make_excel_input(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], "application/zip")
        self.assertEqual(response.headers["x-batch-input-count"], "3")
        self.assertEqual(response.headers["x-batch-paid-count"], "1")
        self.assertEqual(response.headers["x-batch-not-paid-count"], "1")
        self.assertEqual(response.headers["x-batch-unresolved-count"], "1")
        self.assertEqual(response.headers["x-batch-input-mode"], "excel")
        with ZipFile(BytesIO(response.content)) as archive:
            self.assertEqual(len(archive.namelist()), 2)

    def test_text_pdf_is_processed_locally_without_calling_ocr(self) -> None:
        provider = FakeGeminiOcrProvider()
        artifact = process_batch_file(
            "input.pdf", make_text_pdf_input(), self.service, ocr_provider=provider
        )

        self.assertEqual(provider.calls, 0)
        self.assertEqual(artifact.input_mode, "pdf-text")
        self.assertEqual(
            (artifact.paid_count, artifact.not_paid_count, artifact.unresolved_count),
            (1, 1, 1),
        )
        with ZipFile(BytesIO(artifact.archive)) as archive:
            self.assertEqual(
                sorted(archive.namelist()),
                [
                    "input_duoc_tra_luong.xlsx",
                    "input_khong_duoc_tra_luong.xlsx",
                ],
            )

    def test_scanned_pdf_uses_gemini_ocr_text(self) -> None:
        provider = FakeGeminiOcrProvider()
        artifact = process_batch_file(
            "scan.pdf", make_scan_pdf_input(), self.service, ocr_provider=provider
        )

        self.assertEqual(provider.calls, 1)
        self.assertEqual(artifact.input_mode, "gemini-ocr")
        self.assertEqual(artifact.input_count, 3)
        self.assertEqual(
            (artifact.paid_count, artifact.not_paid_count, artifact.unresolved_count),
            (1, 1, 1),
        )

    def test_api_returns_503_when_scanned_pdf_has_no_ocr_configuration(self) -> None:
        with patch("api.main.ocr_provider", UnconfiguredOcrProvider()):
            response = TestClient(app).post(
                "/api/v1/organizations/batch",
                files={"file": ("scan.pdf", make_scan_pdf_input(), "application/pdf")},
            )

        self.assertEqual(response.status_code, 503)
        self.assertIn("Gemini", response.json()["detail"])

    def test_gemini_ocr_calls_generate_content_with_inline_pdf(self) -> None:
        response = MagicMock(status_code=200)
        response.json.return_value = {
            "candidates": [
                {
                    "content": {
                        "parts": [{"text": '["Công an tỉnh Thái Bình"]'}]
                    },
                }
            ]
        }
        with patch("requests.post", return_value=response) as post:
            markdown = GeminiOcrProvider(
                api_url="https://generativelanguage.googleapis.com/v1beta",
                api_key="test-key",
                model="gemini-2.5-flash",
                max_output_tokens=4096,
            ).extract_text(make_scan_pdf_input())

        self.assertIn("Công an tỉnh Thái Bình", markdown)
        self.assertEqual(
            post.call_args.args[0],
            "https://generativelanguage.googleapis.com/v1beta/models/gemini-2.5-flash:generateContent",
        )
        self.assertEqual(post.call_args.kwargs["headers"]["x-goog-api-key"], "test-key")
        self.assertNotIn("test-key", post.call_args.args[0])
        self.assertIs(post.call_args.kwargs["json"]["store"], False)
        self.assertEqual(
            post.call_args.kwargs["json"]["generationConfig"]["maxOutputTokens"],
            4096,
        )
        self.assertEqual(
            post.call_args.kwargs["json"]["generationConfig"]["thinkingConfig"],
            {"thinkingBudget": 0},
        )
        inline_pdf = post.call_args.kwargs["json"]["contents"][0]["parts"][0][
            "inline_data"
        ]
        self.assertEqual(inline_pdf["mime_type"], "application/pdf")
        self.assertTrue(
            __import__("base64").b64decode(inline_pdf["data"]).startswith(b"%PDF-")
        )

    def test_ocr_dns_failure_returns_actionable_message(self) -> None:
        with patch(
            "requests.post",
            side_effect=requests.ConnectionError(
                "NameResolutionError: Failed to resolve 'generativelanguage.googleapis.com'"
            ),
        ):
            with self.assertRaisesRegex(OcrUpstreamError, "hostname"):
                GeminiOcrProvider(
                    api_url="https://generativelanguage.googleapis.com/v1beta",
                    api_key="test-key",
                ).extract_text(make_scan_pdf_input())

    def test_gemini_auth_failure_does_not_expose_upstream_body(self) -> None:
        response = MagicMock(status_code=401, text="secret upstream response")
        with patch("requests.post", return_value=response):
            with self.assertRaisesRegex(OcrUpstreamError, "GEMINI_API_KEY") as raised:
                GeminiOcrProvider(api_key="test-key").extract_text(
                    make_scan_pdf_input()
                )

        self.assertNotIn("secret upstream response", str(raised.exception))

    def test_gemini_ocr_ignores_thoughts_and_accepts_fenced_json(self) -> None:
        response = MagicMock(status_code=200)
        response.json.return_value = {
            "candidates": [
                {
                    "content": {
                        "parts": [
                            {"thought": True, "text": "not JSON"},
                            {"text": '```json\n["Công an tỉnh Thái Bình"]\n```'},
                        ]
                    },
                    "finishReason": "STOP",
                }
            ],
            "responseId": "safe-response-id",
        }
        with patch("requests.post", return_value=response):
            result = GeminiOcrProvider(api_key="test-key").extract_text(
                make_scan_pdf_input()
            )

        self.assertEqual(result, "Công an tỉnh Thái Bình")

    def test_gemini_max_tokens_returns_actionable_error(self) -> None:
        response = MagicMock(status_code=200)
        response.json.return_value = {
            "candidates": [{"finishReason": "MAX_TOKENS"}],
            "responseId": "safe-response-id",
        }
        with patch("requests.post", return_value=response):
            with self.assertRaisesRegex(OcrUpstreamError, "MAX_TOKENS"):
                GeminiOcrProvider(api_key="test-key").extract_text(
                    make_scan_pdf_input()
                )

    def test_gemini_prompt_block_returns_reason(self) -> None:
        response = MagicMock(status_code=200)
        response.json.return_value = {
            "promptFeedback": {"blockReason": "PROHIBITED_CONTENT"},
            "responseId": "safe-response-id",
        }
        with patch("requests.post", return_value=response):
            with self.assertRaisesRegex(OcrUpstreamError, "PROHIBITED_CONTENT"):
                GeminiOcrProvider(api_key="test-key").extract_text(
                    make_scan_pdf_input()
                )


if __name__ == "__main__":
    unittest.main()
